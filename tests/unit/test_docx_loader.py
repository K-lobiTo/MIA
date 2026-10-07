from docx import Document

from mia.ingestion.loaders.docx_loader import DocxLoader


def test_lee_parrafos_y_tablas_en_orden_sin_repetir_celdas_combinadas(tmp_path):
    doc = Document()
    doc.add_paragraph("Programa del curso")
    table = doc.add_table(rows=3, cols=2)
    titulo = table.cell(0, 0).merge(table.cell(0, 1))
    titulo.text = "1 Datos generales"
    table.cell(1, 0).text = "Nombre del curso:"
    table.cell(1, 1).text = "Cibercrimen"
    table.cell(2, 0).text = "Créditos:"
    table.cell(2, 1).add_table(rows=1, cols=1).cell(0, 0).text = "3"
    doc.add_paragraph("Fin del programa")
    path = tmp_path / "programa.docx"
    doc.save(path)

    lines = DocxLoader().load(path).splitlines()

    assert lines[0] == "Programa del curso"
    assert "1 Datos generales" in lines
    assert "1 Datos generales | 1 Datos generales" not in lines
    assert "Nombre del curso: | Cibercrimen" in lines
    assert any(line.startswith("Créditos: |") for line in lines)
    assert "3" in "\n".join(lines)
    assert lines[-1] == "Fin del programa"
